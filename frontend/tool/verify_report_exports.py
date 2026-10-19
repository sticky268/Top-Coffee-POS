"""Validate report fixtures with independent Excel and PDF readers."""
from pathlib import Path
import re
from openpyxl import load_workbook
from pypdf import PdfReader

root = Path("build/report-exports")
book = load_workbook(root / "sales.xlsx", data_only=False)
sheet = book["Sales report"]
assert sheet["A1"].value == "Top Coffee POS sales report"
assert sheet["B2"].value == "Cafe & Coffee"
assert any(c.value == 12.5 and c.data_type == "n" for row in sheet for c in row)
formula_name = next(c for row in sheet for c in row if c.value == '=1+1, "Latte"')
assert formula_name.data_type == "s", "Product names must remain literal text"
reader = PdfReader(root / "sales.pdf")
assert reader.pages
text = "\n".join(page.extract_text() for page in reader.pages)
compact = re.sub(r"\s+", "", text)
assert "TopCoffee" in compact and "12.5" in compact, repr(text)
print("Excel opens correctly, amounts are numeric, formula names are text, and PDF contains report data.")

# New report money stays numeric and the fourth (net) PDF column is retained.
exact = load_workbook(root / "sales-exact.xlsx", data_only=False)["Sales report"]
assert exact["D4"].value == 8.01 and exact["D4"].data_type == "n"
assert exact["B5"].value == -0.1 and exact["B5"].data_type == "n"
exact_text = "".join(page.extract_text() for page in PdfReader(root / "sales-exact.pdf").pages)
assert "8.01" in exact_text and "-0.10" in exact_text, repr(exact_text)
print("Exact-money workbook and four-column PDF preserve net and negative amounts.")

profit = load_workbook(root / "profit-exact.xlsx", data_only=False)["Sales report"]
for value in [1.123457, 4.876543]:
    assert any(c.value == value and c.data_type == "n" for row in profit for c in row)
profit_text = "".join(page.extract_text() for page in PdfReader(root / "profit-exact.pdf").pages)
assert "Profitability" in profit_text and "1.123457" in profit_text and "4.876543" in profit_text
assert "Riverside" in profit_text and "2026-10-19" in profit_text
profit_rows = list(profit.iter_rows())
compact_profit_text = re.sub(r"\s+", "", profit_text)
for group in ["Products", "Categories", "Branches"]:
    start = next(i for i, row in enumerate(profit_rows) if row[0].value == f"{group} comparisons")
    assert [c.value for c in profit_rows[start + 1][:4]] == ["Name", "Metric", "Previous", "Change %"]
    assert profit_rows[start + 2][1].value == "Units sold"
    assert profit_rows[start + 2][2].value == 1 and profit_rows[start + 2][2].data_type == "n"
    assert profit_rows[start + 3][1].value == "Recipe COGS"
    assert profit_rows[start + 3][2].value == 1 and profit_rows[start + 3][2].data_type == "n"
    assert f"{group}comparisons" in compact_profit_text
assert "Change(percentagepoints)" in compact_profit_text
print("Profit exports preserve six-decimal COGS, profit, branch and dates.")

activity = load_workbook(root / "activity-exact.xlsx", data_only=False)["Sales report"]
assert any(c.value == 24.98 and c.data_type == "n" for row in activity for c in row)
cashier = next(c for row in activity for c in row if isinstance(c.value, str) and c.value.startswith("=Cashier"))
assert cashier.data_type == "s"
activity_rows = list(activity.iter_rows())
assert sum(isinstance(row[0].value, str) and row[0].value.endswith("/ Paid orders") for row in activity_rows) == 168
activity_pdf = PdfReader(root / "activity-exact.pdf")
activity_text = re.sub(r"\s+", "", "".join(page.extract_text() for page in activity_pdf.pages))
for expected in ["24.98", "Riverside", "2026-10-19", "Change(points)", "Discount", "Void", "Cashier"]:
    assert expected in activity_text
print("Time/cashier exports preserve exact numeric amounts, literal names, 168 Excel cells and labelled PDF comparisons.")

# Phase 4 inventory fixtures retain numeric decimals and make missing history explicit.
inventory = load_workbook(root / 'inventory-exact.xlsx', data_only=False)['Sales report']
assert any(c.value == 1.999002 and c.data_type == 'n' for row in inventory for c in row)
assert any(c.value == 1.999 and c.data_type == 'n' for row in inventory for c in row)
name = next(c for row in inventory for c in row if c.value == '=Coffee')
assert name.data_type == 's'
assert any(c.value == 'Unavailable' for row in inventory for c in row)
inventory_pdf = PdfReader(root / 'inventory-exact.pdf')
inventory_text = re.sub(r'\s+', '', ''.join(page.extract_text() for page in inventory_pdf.pages))
for expected in ['Inventory', 'Riverside', '2026-10-19', '1.999002', 'Theoretical', 'Unavailable', 'Change%', 'Differentingredientunitsareneversummed']:
    assert expected in inventory_text, expected
print('Inventory exports preserve numeric six-decimal costs, three-decimal quantities, unavailable history and branch/date context.')
